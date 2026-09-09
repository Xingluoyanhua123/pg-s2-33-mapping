from pathlib import Path

from openpyxl import load_workbook

from .models import Course


EXCEL_PATH = (
    Path(__file__).resolve().parents[1]
    / "data"
    / "sample_mapping.xlsx"
)

DEMO_SUGGESTIONS = {
    "INPR140285E": (
        "COMP1002",
        "Programming Fundamentals",
    ),
    "DBSY240184E": (
        "COMP2003",
        "Database Systems",
    ),
}


def clean_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def load_courses() -> list[Course]:
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(
            f"Sample mapping workbook was not found: {EXCEL_PATH}"
        )

    workbook = load_workbook(
        EXCEL_PATH,
        read_only=True,
        data_only=True,
    )
    worksheet = workbook.active

    courses: list[Course] = []

    for row in worksheet.iter_rows(
        min_row=5,
        max_col=15,
        values_only=True,
    ):
        source_code = clean_text(row[0])

        if not source_code:
            continue

        source_name = clean_text(row[1])
        description = clean_text(row[2])
        learning_outcomes = clean_text(row[3])

        try:
            units = int(float(row[4] or 0))
        except (TypeError, ValueError):
            units = 0

        target_code = clean_text(row[5])
        target_name = clean_text(row[6])

        if not target_code and source_code in DEMO_SUGGESTIONS:
            target_code, target_name = DEMO_SUGGESTIONS[source_code]

        if not target_code:
            target_code = "PENDING"

        if not target_name:
            target_name = "Requires university review"

        courses.append(
            Course(
                id=len(courses) + 1,
                code=source_code,
                name=source_name,
                units=units,
                description=description,
                learning_outcomes=learning_outcomes,
                suggested_target_code=target_code,
                suggested_target_name=target_name,
            )
        )

    workbook.close()
    return courses


COURSES = load_courses()